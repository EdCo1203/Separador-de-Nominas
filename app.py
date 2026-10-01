import streamlit as st
import io
import re
import zipfile
import unicodedata
from typing import Optional
from pypdf import PdfReader, PdfWriter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def limpiar_para_archivo(s: str) -> str:
    s = s.strip().upper()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    s = re.sub(r"[^A-Z0-9 _-]", "", s)
    s = re.sub(r"\s+", " ", s)
    s = s.replace(" ", "_")
    return s or "SIN_NOMBRE"


def extraer_nombre(texto: str) -> Optional[str]:
    patron = r"TRABAJADOR/A.*?\n\s*([A-ZÁÉÍÓÚÑ ]{5,})\s+PERSONAL"
    m = re.search(patron, texto, re.DOTALL)
    if m:
        return m.group(1).strip()
    return None


def extraer_dni(texto: str) -> str:
    # El DNI aparece tras la fecha de antigüedad en la línea de TRABAJADOR/A
    # Formato: DD MMM YY seguido del DNI (ej: Y2929822G)
    patron = r"\d{1,2}\s+[A-Z]{3}\s+\d{2}\s+([A-Z0-9]{7,10})\s"
    m = re.search(patron, texto)
    if m:
        return m.group(1).strip()
    return "SIN_DNI"


def extraer_periodo(texto: str) -> str:
    patron = r"MENS\s+(\d{1,2})\s+([A-ZÁÉÍÓÚ]{3})\s+(\d{2})\s+a"
    m = re.search(patron, texto)
    if not m:
        return "0000-00"

    mes_txt, anio = m.group(2), m.group(3)

    meses = {
        "ENE": "01", "FEB": "02", "MAR": "03", "ABR": "04",
        "MAY": "05", "JUN": "06", "JUL": "07", "AGO": "08",
        "SEP": "09", "OCT": "10", "NOV": "11", "DIC": "12"
    }

    return f"20{anio}-{meses.get(mes_txt, '00')}"


def generar_excel(registros: list, periodo: str) -> bytes:
    """Genera un Excel con columnas: Nº, Nombre, DNI, Periodo, Archivo PDF."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nóminas"

    # Cabecera
    headers = ["Nº", "Nombre", "DNI", "Periodo", "Archivo PDF"]
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E4057")
    header_align = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Datos
    row_font = Font(name="Arial", size=11)
    alt_fill = PatternFill("solid", fgColor="EFF3F8")

    for i, reg in enumerate(registros, 1):
        fill = alt_fill if i % 2 == 0 else PatternFill()
        row_data = [i, reg["nombre"], reg["dni"], reg["periodo"], reg["archivo"]]
        for col, valor in enumerate(row_data, 1):
            cell = ws.cell(row=i + 1, column=col, value=valor)
            cell.font = row_font
            cell.fill = fill
            cell.alignment = Alignment(
                horizontal="center" if col in [1, 3, 4] else "left"
            )

    # Anchos de columna
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 45

    ws.row_dimensions[1].height = 20

    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_bytes.seek(0)
    return excel_bytes.read()


# --- Streamlit APP ---
st.title("📄 Separador de Nóminas PDF – Entregalia Tools")
uploaded_file = st.file_uploader("Sube el PDF con las nóminas", type=["pdf"])

if uploaded_file:
    pdf_reader = PdfReader(uploaded_file)
    total_pages = len(pdf_reader.pages)

    st.info(f"📑 El PDF tiene {total_pages} páginas.")

    if st.button("Generar PDFs individuales"):
        progress = st.progress(0)
        status = st.empty()

        primer_texto = (pdf_reader.pages[0].extract_text() or "")
        periodo_zip = extraer_periodo(primer_texto)
        carpeta_zip = f"nominas_{periodo_zip}"

        zip_buffer = io.BytesIO()
        usados = {}
        registros = []

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zipf:
            for i, page in enumerate(pdf_reader.pages):
                status.text(f"Procesando página {i+1} de {total_pages}…")

                texto = page.extract_text() or ""
                nombre = extraer_nombre(texto) or f"Pagina_{i+1}"
                nombre_limpio = limpiar_para_archivo(nombre)
                periodo = extraer_periodo(texto)
                dni = extraer_dni(texto)

                nombre_base = f"{nombre_limpio}_{periodo}"

                count = usados.get(nombre_base, 0) + 1
                usados[nombre_base] = count
                sufijo = f"_{count}" if count > 1 else ""

                filename_pdf = f"{nombre_base}{sufijo}.pdf"
                filepath_zip = f"{carpeta_zip}/{filename_pdf}"

                writer = PdfWriter()
                writer.add_page(page)

                pdf_bytes = io.BytesIO()
                writer.write(pdf_bytes)
                pdf_bytes.seek(0)

                zipf.writestr(filepath_zip, pdf_bytes.read())

                registros.append({
                    "nombre": nombre,
                    "dni": dni,
                    "periodo": periodo,
                    "archivo": filename_pdf
                })

                progress.progress((i + 1) / total_pages)

            # Añadir Excel al ZIP
            excel_bytes = generar_excel(registros, periodo_zip)
            zipf.writestr(
                f"{carpeta_zip}/listado_nominas_{periodo_zip}.xlsx",
                excel_bytes
            )

        zip_buffer.seek(0)
        st.success(f"✔ Proceso finalizado. {total_pages} nóminas procesadas.")
        st.download_button(
            label="⬇️ Descargar ZIP con nóminas separadas",
            data=zip_buffer,
            file_name=f"{carpeta_zip}.zip",
            mime="application/zip"
        )
